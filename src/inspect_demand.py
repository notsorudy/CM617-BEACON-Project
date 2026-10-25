"""BEACON: look at the structure of the Grid-India / Mendeley demand files (prints only).
Reads with openpyxl read_only so the 48 MB files do not load fully into memory."""
import sys
from pathlib import Path

from openpyxl import load_workbook

DIR = Path(__file__).resolve().parents[1] / "data" / "demand" / "Electricity Demand, Solar and Wind Generation Data"
names = sys.argv[1:] or ["September 2021.xlsx", "September 2023.xlsx", "January 2024- June 2025.xlsx"]
for n in names:
    wb = load_workbook(DIR / n, read_only=True, data_only=True)
    print(f"=== {n}: sheets {wb.sheetnames}")
    for ws in wb.worksheets:
        print(f"  sheet {ws.title!r}: dimensions {ws.calculate_dimension()}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            print("   ", row[:14])
            if i >= 8:
                break
    wb.close()
